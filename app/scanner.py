import asyncio
import os
import re
import urllib.parse
import xml.etree.ElementTree as ElementTree
from collections import Counter, deque
from io import BytesIO

import requests
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from playwright.async_api import async_playwright

load_dotenv()

TRACKING_PARAMETERS  = {"fbclid", "gclid", "mc_cid", "mc_eid"}
MAX_RESOURCE_BYTES   = int(os.getenv("MAX_RESOURCE_BYTES", str(25 * 1024 * 1024)))


# ── HTTP helper ───────────────────────────────────────────────────────────────

def safe_get(url: str, timeout: int = 20, headers: dict | None = None):
    hdrs = headers or {"User-Agent": "EMBEDD-IQ crawler"}
    try:
        return requests.get(url, timeout=timeout, headers=hdrs, verify=True)
    except requests.exceptions.SSLError:
        return requests.get(url, timeout=timeout, headers=hdrs, verify=False)


# ── URL helpers ───────────────────────────────────────────────────────────────

def normalize_url(url: str, root_host: str):
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme not in {"http", "https"}:
        return None
    hostname = (parsed.hostname or "").lower().rstrip(".")
    root     = root_host.lower().removeprefix("www.").rstrip(".")
    if hostname.removeprefix("www.") != root:
        return None
    query = [(k, v) for k, v in urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
             if k.lower() not in TRACKING_PARAMETERS and not k.lower().startswith("utm_")]
    path = parsed.path or "/"
    if path != "/":
        path = path.rstrip("/")
    return urllib.parse.urlunsplit((parsed.scheme, hostname, path, urllib.parse.urlencode(query), ""))


def discover_sitemap_urls(start_url: str, root_host: str) -> set:
    origin = "{}://{}".format(urllib.parse.urlsplit(start_url).scheme, root_host)
    pending = deque([
        urllib.parse.urljoin(origin, "/sitemap.xml"),
        urllib.parse.urljoin(origin, "/sitemap_index.xml"),
    ])
    try:
        robots = safe_get(urllib.parse.urljoin(origin, "/robots.txt"), timeout=10)
        if robots.ok:
            pending.extend(re.findall(r"(?im)^\s*sitemap:\s*(\S+)", robots.text))
    except Exception:
        pass

    found, checked = set(), set()
    while pending:
        smap = pending.popleft()
        if smap in checked:
            continue
        checked.add(smap)
        try:
            r = safe_get(smap, timeout=10)
            if r.ok and "xml" in r.headers.get("content-type", "").lower():
                root_el = ElementTree.fromstring(r.content)
                kind    = root_el.tag.rsplit("}", 1)[-1]
                for loc in root_el.iter():
                    if loc.tag.rsplit("}", 1)[-1] == "loc" and loc.text:
                        if kind == "sitemapindex":
                            pending.append(loc.text.strip())
                        else:
                            n = normalize_url(loc.text.strip(), root_host)
                            if n:
                                found.add(n)
        except Exception:
            pass
    return found


# ── Brand extraction ──────────────────────────────────────────────────────────

def _normalize_hex(value: str):
    if not value:
        return None
    raw = value.strip().lower()
    if raw.startswith("#"):
        if re.fullmatch(r"#[0-9a-f]{3}", raw):
            return "#" + "".join(c * 2 for c in raw[1:])
        if re.fullmatch(r"#[0-9a-f]{6}", raw):
            return raw
        return None
    m = re.fullmatch(r"rgba?\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})[^)]*\)", raw)
    if m:
        return "#%02x%02x%02x" % tuple(int(x) for x in m.groups())
    return None


def extract_brand_info(url: str, html: str) -> dict:
    soup  = BeautifulSoup(html, "html.parser")
    title = soup.title.string.strip() if soup.title and soup.title.string else "Chatbot"

    logo_url = None
    
    # 1. Look for explicit logo meta tags
    og_logo = soup.find("meta", property=lambda v: v and v.lower() == "og:logo")
    if og_logo and og_logo.get("content"):
        logo_url = urllib.parse.urljoin(url, og_logo["content"])

    # 2. Look for standard WordPress custom-logo class
    if not logo_url:
        img = soup.find("img", class_=lambda c: c and "custom-logo" in " ".join(c).lower())
        if img:
            src = img.get("src") or img.get("data-src") or img.get("data-lazy-src")
            if src:
                logo_url = urllib.parse.urljoin(url, src)

    # 3. Look for image inside a link to the homepage
    if not logo_url:
        parsed_url = urllib.parse.urlsplit(url)
        root_path = "/"
        root_url = f"{parsed_url.scheme}://{parsed_url.netloc}/"
        root_url_no_slash = f"{parsed_url.scheme}://{parsed_url.netloc}"
        
        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            if href in [root_path, root_url, root_url_no_slash, url]:
                img = a.find("img")
                if img:
                    src = img.get("src") or img.get("data-src") or img.get("data-lazy-src")
                    if src:
                        logo_url = urllib.parse.urljoin(url, src)
                        break

    # 4. Fallback to any img with 'logo' in class/id/alt
    if not logo_url:
        for img in soup.find_all("img"):
            hay = " ".join([img.get("id") or "", " ".join(img.get("class") or []),
                            img.get("alt") or ""]).lower()
            src = img.get("src") or img.get("data-src") or img.get("data-lazy-src")
            if src and "logo" in hay:
                logo_url = urllib.parse.urljoin(url, src)
                break

    # 5. Fallback to og:image
    if not logo_url:
        og_img = soup.find("meta", property=lambda v: v and v.lower() == "og:image")
        if og_img and og_img.get("content"):
            logo_url = urllib.parse.urljoin(url, og_img["content"])

    # 6. Fallback to favicon
    if not logo_url:
        icon = soup.find("link", rel=lambda v: v and "icon" in v.lower())
        if icon and icon.get("href"):
            logo_url = urllib.parse.urljoin(url, icon["href"])

    # Extract Primary Color
    primary = None
    
    # 1. theme-color meta
    theme_meta = soup.find("meta", attrs={"name": lambda x: x and x.lower() == "theme-color"}) or \
                 soup.find("meta", attrs={"property": lambda x: x and x.lower() == "theme-color"})
    if theme_meta and theme_meta.get("content"):
        primary = _normalize_hex(theme_meta["content"])

    # 2. Look in <style> tags for common CSS variables
    if not primary:
        style_text = ""
        for style in soup.find_all("style"):
            style_text += style.get_text(separator=" ") + " "
        
        match = re.search(r'--(?:primary|brand|accent|main)(?:-color)?\s*:\s*(#[0-9a-fA-F]{3,6})', style_text, re.I)
        if match:
            primary = _normalize_hex(match.group(1))

    # 3. Fallback: find all hex colors in HTML, filter out grays, pick most frequent
    if not primary:
        colors = []
        for raw in re.findall(r'#[0-9a-fA-F]{3,6}\b', html):
            h = _normalize_hex(raw)
            if h:
                try:
                    c = h.lstrip('#')
                    if len(c) == 3: c = ''.join([x*2 for x in c])
                    r, g, b = int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
                    # Filter out neutral/gray colors (where r,g,b are very close)
                    if max(r,g,b) - min(r,g,b) > 25:
                        colors.append(h)
                except:
                    pass
        
        if colors:
            counts = Counter(colors)
            primary = max(counts, key=lambda c: (counts[c], c))

    if not primary:
        primary = "#17191d"

    return {"name": title, "logo_url": logo_url, "primary_color": primary}


# ── Text extraction ───────────────────────────────────────────────────────────

_NOISE_TAGS = ["script", "style", "noscript", "svg", "iframe", "header", "footer", "nav", "aside"]
_NOISE_SEL  = [
    "[role='navigation']", "[role='banner']", "[role='contentinfo']", "[role='complementary']",
    ".cookie-banner", ".cookie-notice", ".nav", ".navbar", ".header", ".site-header",
    ".footer", ".site-footer", ".sidebar", ".widget-area", ".breadcrumb",
    ".advertisement", ".ads", ".adsbygoogle",
    "#nav", "#header", "#footer", "#sidebar", "#cookie-banner",
]


def extract_page_text(html: str, page_url: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for sel in _NOISE_SEL:
        for el in soup.select(sel):
            el.decompose()
    for tag in soup(_NOISE_TAGS):
        tag.decompose()
    for img in soup.find_all("img"):
        alt = (img.get("alt") or img.get("title") or "").strip()
        src = img.get("src") or ""
        if src and not src.startswith("data:"):
            full = urllib.parse.urljoin(page_url, src)
            img.replace_with(f" [Image: {alt} ({full})] " if alt else f" [Image: {full}] ")
        else:
            img.decompose()
    root = (
        soup.find("main")
        or soup.find("article")
        or soup.find(id=re.compile(r'(main|content|primary)', re.I))
        or soup.find(class_=re.compile(r'(main|content|primary)', re.I))
        or soup.body or soup
    )
    raw   = root.get_text(separator=" ")
    clean = re.sub(r'[ \t]+', ' ', raw)
    clean = re.sub(r'\n{3,}', '\n\n', clean)
    return clean.strip()


def extract_resource_text(data: bytes, url: str, content_type: str) -> str:
    ext  = os.path.splitext(urllib.parse.urlsplit(url).path.lower())[1]
    ct   = content_type.lower().split(";", 1)[0].strip()
    try:
        if ct == "application/pdf" or ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(data)).pages).strip()
        if ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or ext == ".docx":
            from docx import Document
            return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs).strip()
        if ct == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append(f"Sheet: {ws.title}")
                rows += [" | ".join(str(v) for v in row if v is not None)
                         for row in ws.iter_rows(values_only=True)]
            return "\n".join(rows).strip()
        if ct.startswith("text/") or ext in {".txt", ".md", ".csv", ".json", ".xml"}:
            return data.decode("utf-8", errors="replace").strip()
    except Exception as e:
        print(f"[Scanner] Resource parse error ({url}): {e}")
    return ""


# ── Async crawler ─────────────────────────────────────────────────────────────

async def _async_scan_website(start_url: str, on_page=None):
    parsed    = urllib.parse.urlsplit(start_url)
    root_host = parsed.hostname or parsed.netloc
    first_url = normalize_url(start_url, root_host)

    pages_data: list = []
    visited: set     = set()
    queue            = asyncio.Queue()

    await queue.put(first_url)
    visited.add(first_url)

    # Pre-populate from sitemap (fast, before browser starts)
    for u in discover_sitemap_urls(first_url, root_host):
        if u not in visited:
            visited.add(u)
            await queue.put(u)

    brand_info: dict = {}
    brand_lock       = asyncio.Lock()
    pages_lock       = asyncio.Lock()

    max_pages    = int(os.getenv("MAX_SCAN_PAGES",    "0"))
    page_timeout = int(os.getenv("SCAN_PAGE_TIMEOUT_MS", "30000"))
    concurrency  = int(os.getenv("SCAN_CONCURRENCY",  "3"))

    # ── Worker drain timeout: wait much longer so slow sites don't starve workers
    QUEUE_WAIT = 15.0   # seconds to wait for the next URL before giving up

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True)
        ctx = await browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/124.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1920, "height": 1080},
            ignore_https_errors=True,
        )
        # Removed resource blocking for stability

        # ── Homepage — scan first to guarantee brand info and seed queue ──────
        try:
            hp   = await ctx.new_page()
            resp = await hp.goto(first_url, timeout=page_timeout, wait_until="commit")
            if resp and "text/html" in resp.headers.get("content-type", "").lower():
                await hp.wait_for_timeout(1500)
                html = await hp.content()
                brand_info = extract_brand_info(first_url, html)
                # Enqueue all links found on homepage
                lsoup = BeautifulSoup(html, "html.parser")
                for a in lsoup.find_all("a", href=True):
                    full = normalize_url(urllib.parse.urljoin(first_url, a["href"]), root_host)
                    if full and full not in visited:
                        visited.add(full)
                        await queue.put(full)
                text = extract_page_text(html, first_url)
                if text:
                    pages_data.append({"url": first_url, "text": text})
                    if on_page:
                        on_page(len(pages_data), first_url)
            await hp.close()
        except Exception as e:
            print(f"[Scanner] Homepage error {first_url}: {e}")

        # ── Workers ───────────────────────────────────────────────────────────
        async def worker():
            nonlocal brand_info
            page = await ctx.new_page()
            try:
                while True:
                    async with pages_lock:
                        if max_pages and len(pages_data) >= max_pages:
                            break
                    # Wait longer for slow sites — don't exit too early
                    try:
                        url = await asyncio.wait_for(queue.get(), timeout=QUEUE_WAIT)
                    except asyncio.TimeoutError:
                        break

                    # Skip if already done
                    if any(p["url"] == url for p in pages_data):
                        queue.task_done()
                        continue

                    try:
                        resp = await page.goto(url, timeout=page_timeout, wait_until="commit")
                        if not resp:
                            queue.task_done()
                            continue

                        ct = resp.headers.get("content-type", "").lower()

                        # Non-HTML resource (PDF, DOCX, etc.)
                        if "text/html" not in ct:
                            try:
                                body = await asyncio.to_thread(
                                    lambda: safe_get(url, timeout=20).content
                                )
                                rtxt = extract_resource_text(body, url, ct)
                                if rtxt:
                                    async with pages_lock:
                                        if not max_pages or len(pages_data) < max_pages:
                                            pages_data.append({"url": url, "text": rtxt})
                                            if on_page:
                                                on_page(len(pages_data), url)
                            except Exception:
                                pass
                            queue.task_done()
                            continue

                        await page.wait_for_timeout(500)
                        html = await page.content()

                        async with brand_lock:
                            if not brand_info:
                                brand_info = extract_brand_info(url, html)

                        # Enqueue new links
                        lsoup = BeautifulSoup(html, "html.parser")
                        for a in lsoup.find_all("a", href=True):
                            full = normalize_url(urllib.parse.urljoin(url, a["href"]), root_host)
                            if full and full not in visited:
                                async with pages_lock:
                                    if not max_pages or len(visited) < max_pages * 3:
                                        visited.add(full)
                                        await queue.put(full)

                        text = extract_page_text(html, url)
                        if text:
                            async with pages_lock:
                                if not max_pages or len(pages_data) < max_pages:
                                    pages_data.append({"url": url, "text": text})
                                    if on_page:
                                        on_page(len(pages_data), url)
                    except Exception as e:
                        print(f"[Scanner] Error {url}: {e}")
                    finally:
                        queue.task_done()
            finally:
                await page.close()

        await asyncio.gather(*[asyncio.create_task(worker()) for _ in range(concurrency)])
        await browser.close()

    return brand_info, pages_data


def scan_website(start_url: str, on_page=None):
    return asyncio.run(_async_scan_website(start_url, on_page=on_page))
